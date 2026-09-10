import sys
sys.path.append('local-agent')
import document_tools
import os
kwargs = {
    'path': 'test_fallback.pptx',
    'title': 'Test',
    'slides': [{'title': 'Fail Download', 'content': 'Will Fail', 'type': 'image', 'image_path': ''}],
    'auto_images': True
}
# Monkey patch to force failure
document_tools.get_image_for_slide = lambda x: 'http://localhost:1/fake.png'
print(document_tools.create_ppt(kwargs))

# Read with pptx to confirm shape
import pptx
prs = pptx.Presentation('test_fallback.pptx')
slide = prs.slides[0]
print(f'Slide count: {len(prs.slides)}')
print(f'Shapes on slide 0: {len(slide.shapes)}')
for shape in slide.shapes:
    print(f'Shape type: {shape.shape_type}')
