import base64
import os
import tempfile
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Simple lightweight SVG icons Base64 encoded
SVG_ICONS = {
    "quote": "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0iY3VycmVudENvbG9yIj48cGF0aCBkPSJNMTQuMDE3IDhsLTIuNzcuMWMtMS45OC4wNy0zLjgxIDEuNjctMy44MSAzLjc0djUuMTZjMCAxLjE1Ljk1IDIuMDkgMi4xIDIuMDloMy43NGMxLjE1IDAgMi4xLS45NCAyLjEtMi4wOXYtMy4xMWMwLTEuMTUtLjk1LTIuMDktMi4xLTIuMDhoLTEuNWMtLjA2LTEuMS43MS0yIDEuOC0yLjFoLjE4bDQuMDEtLjE1di0yLjExbC00LjAxLS4xNnptLTguMzMgMGwtMi43Ny4xYy0xLjk4LjA3LTMuODEgMS42Ny0zLjgxIDMuNzR2NS4xNmMwIDEuMTUuOTUgMi4wOSAyLjEgMi4wOWgzLjc0YzEuMTUgMCAyLjEtLjk0IDIuMS0yLjA5di0zLjExYzAtMS4xNS0uOTUtMi4wOS0yLjEtMi4wOGgtMS41Yy0uMDYtMS4xLjcxLTIgMS44LTIuMWguMThsNC4wMS0uMTV2LTIuMTFsLTQuMDEtLjE2eiIvPjwvc3ZnPg==",
    "chart": "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0iY3VycmVudENvbG9yIj48cGF0aCBkPSJNNyAyNGgtNnYtMTRoNnYxNHptOC0xOGgtNnYxOGg2di0xOHptOCAxMGgtNnY4aDZ2LTh6Ii8+PC9zdmc+",
    "check": "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0iY3VycmVudENvbG9yIj48cGF0aCBkPSJNMjAuMjggNS45N2wtMTEuOTkgMTEuOTktNS41Ny01LjU3LTEuNDEgMS40MSA2Ljk4IDYuOTggMTMuNDEtMTMuNDF6Ii8+PC9zdmc+",
    "cloud": "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0iY3VycmVudENvbG9yIj48cGF0aCBkPSJNMTcuNSAxOS4yNWgtMTAuNWMtMi44OSAwLTUuMjUtMi4zNi01LjI1LTUuMjVzMi4zNi01LjI1IDUuMjUtNS4yNWMuMTEgMCAuMjIgMCAuMzMuMDFjLjgzLTMuODUgNC4yNy02Ljc2IDguNDItNi43NiA0Ljc5IDAgOC42NyAzLjg4IDguNjcgOC42NyAwIDEuNS0uMzggMi45Mi0xLjA3IDQuMTYgMS40OS43OSAyLjU3IDIuMzggMi41NyA0LjE3IDAgMi42NS0yLjE1IDQuOC00LjggNC44em0tMTAuNS04LTVjMS43OSAwIDMuMjUgMS40NiAzLjI1IDMuMjVzLTEuNDYgMy4yNS0zLjI1IDMuMjVoMTAuNWMxLjUxIDAgMi43NS0xLjIzIDIuNzUtMi43NXMtMS4yMy0yLjc1LTIuNzUtMi43NWgtLjc1di0xYy0xLjU1IDAtMi44MS0xLjI2LTIuODEtMi44MXMyLjgtMi44MSA2LjI2LTIuODFjMy40NiAwIDYuMjYgMi44MSA2LjI2IDYuMjZ6Ii8+PC9zdmc+",
    "user": "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0iY3VycmVudENvbG9yIj48cGF0aCBkPSJNMTIgMmMtMi43NiAwLTUgMi4yNC01IDVzMi4yNCA1IDUgNSA1LTIuMjQgNS01LTIuMjQtNS01LTV6bTAgOGMtMS42NiAwLTMtMS4zNC0zLTNzMS4zNC0zIDMtMyAzIDEuMzQgMyAzLTEuMzQgMy0zIDN6bS03LjUgMTAuN2MyLjYtMi44MiA1LjktNC43IDkuNS00LjcgMy42IDAgNi45IDEuODggOS41IDQuN3YxLjNoLTE5di0xLjN6bTItMWMtMi4wNS0yLjUtNC43My00LTcuNS00LTIuNzcgMC01LjQ1IDEuNS03LjUgNHoiLz48L3N2Zz4="
}

def apply_background(slide, theme, prs):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*theme['color_primary'])
    bg.line.fill.background()

def add_card(slide, x, y, cx, cy, fill_color, border_color):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cx, cy)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(*fill_color)
    card.line.color.rgb = RGBColor(*border_color)
    card.line.width = Pt(1)
    return card

def set_run_font(run, theme, role="Body"):
    if role == "Title":
        run.font.name = theme.get("font_heading", "Arial")
        run.font.color.rgb = RGBColor(*theme['color_heading'])
    elif role == "Subtitle":
        run.font.name = theme.get("font_body", "Arial")
        run.font.color.rgb = RGBColor(*theme['color_accent'])
    else:
        run.font.name = theme.get("font_body", "Arial")
        run.font.color.rgb = RGBColor(*theme['color_text'])

def add_footer(slide, title, page_idx, page_count, theme):
    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(9), Inches(0.4))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"{title}   |   {page_idx}/{page_count}"
    set_run_font(p.runs[0], theme, "Body")
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(*theme['color_accent'])
    p.alignment = PP_ALIGN.RIGHT

def get_svg_icon_path(icon_name):
    if icon_name not in SVG_ICONS:
        return None
    svg_data = base64.b64decode(SVG_ICONS[icon_name]).decode('utf-8')
    fd, temp_svg = tempfile.mkstemp(suffix=".svg")
    with os.fdopen(fd, 'w', encoding='utf-8') as f:
        f.write(svg_data)
    import diagram_render
    png_path = diagram_render.svg_file_to_png(temp_svg)
    return png_path
