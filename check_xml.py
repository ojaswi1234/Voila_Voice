import pptx
from lxml import etree
from pptx.oxml.ns import qn

prs = pptx.Presentation('test_agentic_ai_cards.pptx')

# Slide 7 = content slide (index 6, 0-based)
slide = prs.slides[6]
spTree = slide.shapes._spTree
shapes = list(spTree)

print(f"Slide 7 raw XML element count: {len(shapes)}")
for i, elem in enumerate(shapes):
    tag = elem.tag.split('}')[-1]
    xml_str = etree.tostring(elem, pretty_print=False).decode()
    has_round = 'roundRect' in xml_str
    has_shadow = 'outerShdw' in xml_str
    print(f"  [{i}] tag={tag:6s} | rounded={has_round} | shadow={has_shadow}")

# Check gradient background on slide 1 (illustrated_light)
print("\n--- Gradient check on slide 1 ---")
slide1 = prs.slides[0]
cSld = slide1._element.find(qn('p:cSld'))
bg = cSld.find(qn('p:bg')) if cSld is not None else None
if bg is not None:
    print("Found p:bg element")
    print(etree.tostring(bg, pretty_print=True).decode()[:400])
else:
    print("No p:bg element (solid fill or not found)")
