import pptx
from pptx.util import Inches
prs = pptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
svg = '<?xml version="1.0" encoding="UTF-8"?><svg width="100" height="100"><rect width="100" height="100" fill="red"/></svg>'
with open('test.svg', 'w') as f: f.write(svg)
try:
    slide.shapes.add_picture('test.svg', Inches(1), Inches(1))
    print('SUCCESS')
except Exception as e:
    print('FAIL:', type(e), e)
